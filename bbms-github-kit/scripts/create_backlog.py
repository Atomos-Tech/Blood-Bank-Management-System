#!/usr/bin/env python3
"""Create the SE mini-project backlog on GitHub from BBMS_Product_Backlog.xlsx.

What it creates (idempotent where GitHub allows it):
  * labels (FR / NFR / user-story / per-module / sp:N)
  * milestones "Sprint 0 (dry run)", "Sprint 1", "Sprint 2" with due dates
  * one issue per backlog item, with story, acceptance criteria, SRS ID,
    story-point label, milestone and assignee
  * a GitHub Project (v2) linked to the repo, with "Story Points" (number)
    and "Sprint" (single select) fields filled in for every issue

Prerequisites (run once on your own machine):
  gh auth login
  gh auth refresh -s project        # Projects v2 needs the 'project' scope

Usage:
  pip install openpyxl
  python3 scripts/create_backlog.py --repo OWNER/REPO --xlsx BBMS_Product_Backlog.xlsx           # dry run
  python3 scripts/create_backlog.py --repo OWNER/REPO --xlsx BBMS_Product_Backlog.xlsx --apply   # create

The sheet is the source of truth: edit it with the team, then run this once.
Re-running is safe: existing issues (matched by title), labels, milestones and
the project are reused, and their Story Points / Sprint fields are refreshed.
"""
import argparse
import json
import subprocess
import sys
from datetime import date, datetime

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

PROJECT_TITLE = "BBMS - SE Mini Project"


def gh(*args, capture=True, check=True):
    result = subprocess.run(["gh", *args], text=True, capture_output=capture)
    if check and result.returncode != 0:
        sys.exit(f"gh {' '.join(args)}\n{result.stderr.strip()}")
    return result.stdout.strip() if capture else ""


def gh_json(*args):
    out = gh(*args)
    return json.loads(out) if out else None


def _d(v):
    return (v.date() if isinstance(v, datetime) else v if isinstance(v, date) else date.fromisoformat(str(v))).isoformat()


def load_sheet(path):
    wb = load_workbook(path, data_only=False)
    plan = wb["Sprint Plan"]
    sprints, team, mode = [], {}, "sprints"
    for row in plan.iter_rows(min_row=3, values_only=True):
        a = row[0]
        if mode == "sprints":
            if a == "Total":
                mode = "seek-team"
            elif a:
                sprints.append({"name": str(a).strip(), "start": _d(row[1]), "due": _d(row[2]), "description": row[5] or ""})
        elif mode == "seek-team":
            if a == "Member":
                mode = "team"
        elif mode == "team":
            if a in (None, "Unassigned"):
                break
            user = (row[1] or "").strip().lstrip("@")
            team[str(a).strip()] = "" if user in ("", "github-username") else user

    bl = wb["Backlog"]
    rows = list(bl.iter_rows(values_only=True))
    hdr = {h: i for i, h in enumerate(rows[0]) if h}
    items = []
    for r in rows[1:]:
        if not r[hdr["Key"]] or not r[hdr["Title"]]:
            continue
        who = (r[hdr["Assignee"]] or "").strip()
        items.append({
            "key": str(r[hdr["Key"]]).strip(),
            "srs": str(r[hdr["SRS ID"]]).strip(),
            "module": (r[hdr["Module"]] or "").strip(),
            "title": str(r[hdr["Title"]]).strip(),
            "story": r[hdr["User Story"]] or "",
            "acceptance": [l.lstrip("-• ").strip() for l in str(r[hdr["Acceptance Criteria"]] or "").splitlines() if l.strip()],
            "points": int(r[hdr["Story Points"]] or 0),
            "priority": r[hdr["Priority"]] or "",
            "sprint": str(r[hdr["Sprint"]]).strip(),
            "member": who,
            "assignee": team.get(who, ""),
            "status": r[hdr["Codebase Status"]] or "",
            "notes": r[hdr["Gap / Notes"]] or "",
        })
    unknown = {i["sprint"] for i in items} - {s["name"] for s in sprints}
    if unknown:
        sys.exit(f"Backlog uses sprint names not in 'Sprint Plan': {unknown}")
    return sprints, items


def issue_title(item):
    return f"{item['key']} [{item['srs']}] {item['title']}"


def issue_body(item):
    criteria = "\n".join(f"- [ ] {c}" for c in item["acceptance"])
    notes = f"\n### Current state in codebase\n{item['status']}. {item['notes']}\n" if item["status"] else ""
    return (
        f"**SRS requirement:** {item['srs']}  \n"
        f"**Story points:** {item['points']}  \n"
        f"**Priority:** {item['priority']}  \n"
        f"**Sprint:** {item['sprint']}\n\n"
        f"### User story\n{item['story']}\n\n"
        f"### Acceptance criteria\n{criteria}\n"
        f"{notes}\n"
        f"_Branch:_ `{item['key']}-short-name` · _PR must say_ `Closes #<this issue>`\n"
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True, help="OWNER/REPO")
    ap.add_argument("--xlsx", default="BBMS_Product_Backlog.xlsx", help="path to the backlog workbook")
    ap.add_argument("--apply", action="store_true", help="create things (default is a dry run)")
    args = ap.parse_args()
    owner, _ = args.repo.split("/", 1)

    sprints, items = load_sheet(args.xlsx)

    print(f"{len(items)} backlog items, {sum(i['points'] for i in items)} story points")
    for s in sprints:
        pts = sum(i["points"] for i in items if i["sprint"] == s["name"])
        print(f"  {s['name']:<20} due {s['due']}  {pts} pts")
    if not args.apply:
        for i in items:
            who = f"@{i['assignee']}" if i["assignee"] else f"{i['member']} (no GitHub username yet - will be unassigned)"
            print(f"  {issue_title(i)}  ({i['points']} pts, {i['sprint']}, {who})")
        print("\nDry run only. Re-run with --apply to create on GitHub.")
        return

    # Labels
    labels = {"user-story": "1d76db", "FR": "0e8a16", "NFR": "5319e7", "documentation": "0075ca"}
    for i in items:
        labels.setdefault(f"sp:{i['points']}", "fbca04")
        if i["module"]:
            labels.setdefault(i["module"].lower(), "c5def5")
    for name, color in labels.items():
        gh("label", "create", name, "--color", color, "--force", "-R", args.repo)
    print(f"labels: {len(labels)} ok")

    # Milestones = sprints
    existing = {m["title"] for m in gh_json("api", f"repos/{args.repo}/milestones?state=all&per_page=100")}
    for s in sprints:
        if s["name"] not in existing:
            gh("api", f"repos/{args.repo}/milestones", "-f", f"title={s['name']}",
               "-f", f"description={s['description']}", "-f", f"due_on={s['due']}T23:59:59Z")
    print(f"milestones: {len(sprints)} ok")

    # Project v2
    projects = gh_json("project", "list", "--owner", owner, "--format", "json")["projects"]
    proj = next((p for p in projects if p["title"] == PROJECT_TITLE), None)
    if not proj:
        proj = gh_json("project", "create", "--owner", owner, "--title", PROJECT_TITLE, "--format", "json")
    number, project_id = str(proj["number"]), proj["id"]
    gh("project", "link", number, "--owner", owner, "--repo", args.repo, check=False)

    fields = {f["name"]: f for f in gh_json("project", "field-list", number, "--owner", owner, "--format", "json")["fields"]}
    if "Story Points" not in fields:
        gh("project", "field-create", number, "--owner", owner, "--name", "Story Points", "--data-type", "NUMBER")
    if "Sprint" not in fields:
        gh("project", "field-create", number, "--owner", owner, "--name", "Sprint", "--data-type", "SINGLE_SELECT",
           "--single-select-options", ",".join(s["name"] for s in sprints))
    fields = {f["name"]: f for f in gh_json("project", "field-list", number, "--owner", owner, "--format", "json")["fields"]}
    sp_field, sprint_field = fields["Story Points"], fields["Sprint"]
    sprint_opts = {o["name"]: o["id"] for o in sprint_field["options"]}
    print(f"project #{number}: {proj.get('url', '')}")

    # Issues
    open_titles = {i["title"]: i["url"] for i in gh_json(
        "issue", "list", "-R", args.repo, "--state", "all", "--limit", "500", "--json", "title,url")}
    for i in items:
        title = issue_title(i)
        url = open_titles.get(title)
        if not url:
            cmd = ["issue", "create", "-R", args.repo, "--title", title, "--body", issue_body(i),
                   "--milestone", i["sprint"], "--label", "user-story",
                   "--label", "NFR" if i["srs"].upper().startswith("NFR") else "FR" if i["srs"].upper().startswith("FR") else "documentation",
                   "--label", f"sp:{i['points']}"]
            if i["module"]:
                cmd += ["--label", i["module"].lower()]
            if i.get("assignee"):
                cmd += ["--assignee", i["assignee"]]
            url = gh(*cmd)
        item = gh_json("project", "item-add", number, "--owner", owner, "--url", url, "--format", "json")
        gh("project", "item-edit", "--id", item["id"], "--project-id", project_id,
           "--field-id", sp_field["id"], "--number", str(i["points"]))
        gh("project", "item-edit", "--id", item["id"], "--project-id", project_id,
           "--field-id", sprint_field["id"], "--single-select-option-id", sprint_opts[i["sprint"]])
        print(f"  ok  {title}  -> {url}")

    print("\nDone. In the project, add a Board view grouped by Status and a Table view grouped by Sprint.")


if __name__ == "__main__":
    main()
